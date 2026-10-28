"""Read-oriented XLSX document and handler built on openpyxl."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, ClassVar

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from anydoc.core.errors import CorruptFileError
from anydoc.core.model import Capability, Metadata, Part
from anydoc.core.registry import FormatHandler

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class SheetInfo:
    """A worksheet: ``dimensions`` is the used range (e.g. ``A1:E4``) or ``None`` if empty."""

    index: int
    name: str
    state: str
    dimensions: str | None
    max_row: int
    max_column: int
    merged_ranges: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class CellInfo:
    """A non-empty cell.

    ``type`` is one of ``string``, ``number``, ``boolean``, ``date``, ``formula``, ``error``.
    For formulas ``formula`` holds the expression and ``value`` the cached result (or ``None``
    if the file carries no cached value); otherwise ``formula`` is ``None``.
    """

    coordinate: str
    row: int
    column: int
    value: Any
    type: str
    formula: str | None = None
    number_format: str = "General"
    bold: bool = False
    italic: bool = False
    font_name: str | None = None
    font_size: float | None = None
    fill_color: str | None = None


@dataclass(frozen=True)
class DefinedNameInfo:
    """A defined name; ``scope`` is the sheet name or ``None`` for workbook scope."""

    name: str
    value: str
    scope: str | None = None


def _cell_type(cell: Cell) -> str:
    if cell.data_type == "f":
        return "formula"
    if cell.data_type == "e":
        return "error"
    if cell.data_type == "b":
        return "boolean"
    if isinstance(cell.value, datetime | date | time | timedelta):
        return "date"
    if cell.data_type == "n":
        return "number"
    return "string"


def _color(color: Any) -> str | None:
    rgb = getattr(color, "rgb", None)
    return rgb if isinstance(rgb, str) else None


def _fill_color(cell: Cell) -> str | None:
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return None
    return _color(fill.fgColor)


def _text_of(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return str(value)


class XlsxDocument:
    """An opened XLSX workbook; ``save`` rewrites it via openpyxl (formulas preserved)."""

    format = "xlsx"

    def __init__(self, wb: Any, cached: Any, *, path: Path | None = None) -> None:
        self._wb = wb
        self._cached = cached
        self.path = path
        self.metadata = self._build_metadata()

    @property
    def native(self) -> Any:
        """The underlying ``openpyxl.Workbook`` (formulas, not cached values)."""
        return self._wb

    # -- structure ---------------------------------------------------------

    def _sheet(self, sheet: int | str) -> Worksheet:
        try:
            ws: Worksheet = (
                self._wb.worksheets[sheet] if isinstance(sheet, int) else self._wb[sheet]
            )
        except (KeyError, IndexError):
            raise KeyError(f"No such sheet: {sheet!r}") from None
        return ws

    def _cached_sheet(self, ws: Worksheet) -> Worksheet:
        cached: Worksheet = self._cached[ws.title]
        return cached

    @staticmethod
    def _is_empty(ws: Worksheet) -> bool:
        return not any(c.value is not None for row in ws.iter_rows() for c in row)

    def sheets(self) -> list[SheetInfo]:
        """Return every worksheet with state, used range and merged ranges."""
        result: list[SheetInfo] = []
        for i, ws in enumerate(self._wb.worksheets):
            empty = self._is_empty(ws)
            result.append(
                SheetInfo(
                    index=i,
                    name=ws.title,
                    state=ws.sheet_state,
                    dimensions=None if empty else ws.dimensions,
                    max_row=0 if empty else ws.max_row,
                    max_column=0 if empty else ws.max_column,
                    merged_ranges=sorted(str(r) for r in ws.merged_cells.ranges),
                )
            )
        return result

    def cells(self, sheet: int | str = 0) -> list[CellInfo]:
        """Return non-empty cells of a sheet in row-major order, with formulas and styles."""
        ws = self._sheet(sheet)
        cached_ws = self._cached_sheet(ws)
        out: list[CellInfo] = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                kind = _cell_type(cell)
                formula = value = None
                if kind == "formula":
                    formula = str(cell.value)
                    value = cached_ws[cell.coordinate].value
                else:
                    value = cell.value
                out.append(
                    CellInfo(
                        coordinate=cell.coordinate,
                        row=cell.row,
                        column=cell.column,
                        value=value,
                        type=kind,
                        formula=formula,
                        number_format=cell.number_format,
                        bold=bool(cell.font.b),
                        italic=bool(cell.font.i),
                        font_name=cell.font.name,
                        font_size=float(cell.font.sz) if cell.font.sz is not None else None,
                        fill_color=_fill_color(cell),
                    )
                )
        return out

    def rows(self, sheet: int | str = 0) -> list[list[Any]]:
        """Return the used range as rows of values (formulas replaced by cached values)."""
        ws = self._cached_sheet(self._sheet(sheet))
        if self._is_empty(ws):
            return []
        return [list(r) for r in ws.iter_rows(values_only=True)]

    def merged_cells(self, sheet: int | str = 0) -> list[str]:
        """Return merged ranges (e.g. ``A1:B2``) of a sheet."""
        return sorted(str(r) for r in self._sheet(sheet).merged_cells.ranges)

    def style_summary(self, sheet: int | str = 0) -> dict[str, Any]:
        """Summarize styling: number formats, fonts, fills and bold/italic counts."""
        formats: dict[str, int] = {}
        fonts: dict[str, int] = {}
        fills: dict[str, int] = {}
        bold = italic = 0
        for c in self.cells(sheet):
            formats[c.number_format] = formats.get(c.number_format, 0) + 1
            if c.font_name:
                fonts[c.font_name] = fonts.get(c.font_name, 0) + 1
            if c.fill_color:
                fills[c.fill_color] = fills.get(c.fill_color, 0) + 1
            bold += c.bold
            italic += c.italic
        return {
            "number_formats": formats,
            "fonts": fonts,
            "fills": fills,
            "bold_cells": bold,
            "italic_cells": italic,
        }

    def defined_names(self) -> list[DefinedNameInfo]:
        """Return workbook- and sheet-scoped defined names."""
        names = [DefinedNameInfo(n, d.attr_text or "") for n, d in self._wb.defined_names.items()]
        for ws in self._wb.worksheets:
            names += [
                DefinedNameInfo(n, d.attr_text or "", ws.title) for n, d in ws.defined_names.items()
            ]
        return names

    # -- Document protocol -------------------------------------------------

    def _build_metadata(self) -> Metadata:
        props = self._wb.properties
        sheets = self.sheets()
        formula_count = sum(
            1 for ws in self._wb.worksheets for r in ws.iter_rows() for c in r if c.data_type == "f"
        )
        return Metadata(
            format="xlsx",
            mime=_XLSX_MIME,
            title=props.title or None,
            author=props.creator or None,
            created=props.created,
            modified=props.modified,
            page_count=None,
            extra={
                "subject": props.subject or None,
                "keywords": props.keywords or None,
                "last_modified_by": props.lastModifiedBy or None,
                "sheet_count": len(sheets),
                "sheet_names": [s.name for s in sheets],
                "formula_count": formula_count,
                "defined_names": [n.name for n in self.defined_names()],
            },
        )

    def parts(self) -> list[Part]:
        """Return one part per worksheet."""
        return [Part(kind="sheet", index=s.index, name=s.name) for s in self.sheets()]

    def text(self) -> str:
        """Return each sheet as its name followed by tab-separated rows (cached values)."""
        blocks: list[str] = []
        for s in self.sheets():
            lines = ["\t".join(_text_of(v) for v in row) for row in self.rows(s.index)]
            blocks.append("\n".join([s.name, *lines]))
        return "\n\n".join(blocks)

    def save(self, path: Path | None = None) -> Path:
        """Write the workbook to ``path`` (default: the source path) and return it."""
        target = path or self.path
        if target is None:
            raise ValueError("No path given and document has no source path")
        self._wb.save(str(target))
        return target

    def close(self) -> None:
        """Close the underlying workbooks."""
        self._wb.close()
        self._cached.close()


class XlsxHandler(FormatHandler):
    """Handler for Excel ``.xlsx`` workbooks."""

    format_id: ClassVar[str] = "xlsx"
    extensions: ClassVar[tuple[str, ...]] = ("xlsx",)
    mime_types: ClassVar[tuple[str, ...]] = (_XLSX_MIME,)
    capabilities: ClassVar[frozenset[Capability]] = frozenset(
        {Capability.READ, Capability.EXTRACT_TEXT}
    )

    @classmethod
    def sniff(cls, head: bytes, path: Path | None) -> float:
        """Return confidence from the extension (content detection lives in ``detect``)."""
        if path is not None and path.suffix.lower() == ".xlsx":
            return 0.9
        return 0.0

    def open(self, source: Path | bytes) -> XlsxDocument:
        """Open an XLSX from a path or bytes; invalid packages raise ``CorruptFileError``."""
        path = source if isinstance(source, Path) else None
        if isinstance(source, Path):
            try:
                data = source.read_bytes()
            except OSError as exc:
                raise CorruptFileError(f"Cannot read {source}: {exc}") from exc
        else:
            data = source
        try:
            wb = load_workbook(io.BytesIO(data))
            cached = load_workbook(io.BytesIO(data), data_only=True)
            return XlsxDocument(wb, cached, path=path)
        except Exception as exc:  # openpyxl raises many types for bad packages
            raise CorruptFileError(f"Invalid XLSX: {exc}") from exc
