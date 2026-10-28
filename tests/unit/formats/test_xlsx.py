"""Tests for the XLSX reader (T013)."""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

import anydoc
from anydoc.core.errors import CorruptFileError
from anydoc.core.model import Capability, Document
from anydoc.core.registry import registry
from anydoc.formats.xlsx import XlsxDocument, XlsxHandler
from tests.fixtures.generate import BODY, TABLE_HEADER, TABLE_ROWS, TITLE

Sample = Callable[[str], Path]


@pytest.fixture(autouse=True)
def _handler() -> None:
    registry.register(XlsxHandler, replace=True)


def _with_cached_value(data: bytes, coordinate: str, value: str) -> bytes:
    """Inject a cached ``<v>`` result for a formula cell (as Excel would write)."""
    out = io.BytesIO()
    pattern = re.compile(rf'(<c r="{coordinate}"[^>]*>\s*<f>[^<]*</f>)\s*(<v\s*/>|<v></v>)?')
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            blob = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text, n = pattern.subn(rf"\1<v>{value}</v>", blob.decode("utf-8"))
                assert n == 1
                blob = text.encode("utf-8")
            zout.writestr(item, blob)
    return out.getvalue()


def _rich_xlsx(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "資料"
    ws["A1"] = "標題"
    ws["A1"].font = Font(bold=True, italic=True, name="Arial", size=14)
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws.merge_cells("A1:C1")
    ws["A2"] = 12.5
    ws["A2"].number_format = "0.00"
    ws["B2"] = True
    ws["C2"] = date(2026, 1, 2)
    ws["D2"] = datetime(2026, 1, 2, 3, 4, 5)
    ws["A3"] = "=SUM(A2:A2)"
    ws["B3"] = "=1/0"
    ws["B3"].data_type = "e"
    hidden = wb.create_sheet("隱藏")
    hidden["A1"] = "secret"
    hidden.sheet_state = "hidden"
    wb.create_sheet("空白")
    wb.defined_names["Total"] = DefinedName("Total", attr_text="資料!$A$3")
    hidden.defined_names["Local"] = DefinedName("Local", attr_text="隱藏!$A$1")
    wb.properties.title = "豐富活頁簿"
    wb.properties.creator = "劉老師"
    buf = io.BytesIO()
    wb.save(buf)
    path = tmp_path / "rich.xlsx"
    path.write_bytes(_with_cached_value(buf.getvalue(), "A3", "12.5"))
    return path


def test_entry_point_and_capabilities() -> None:
    from anydoc.core.registry import HandlerRegistry

    reg = HandlerRegistry()
    reg.load_entry_points()
    assert reg.get("xlsx") is XlsxHandler
    assert Capability.READ in XlsxHandler.capabilities
    assert Capability.EXTRACT_TEXT in XlsxHandler.capabilities


def test_sniff() -> None:
    assert XlsxHandler.sniff(b"", Path("a.XLSX")) == 0.9
    assert XlsxHandler.sniff(b"", Path("a.xls")) == 0.0
    assert XlsxHandler.sniff(b"", None) == 0.0


def test_open_sample_via_public_api(sample: Sample) -> None:
    doc = anydoc.open(sample("xlsx"))
    assert isinstance(doc, XlsxDocument)
    assert isinstance(doc, Document)
    assert doc.format == "xlsx"
    assert doc.metadata.title == TITLE
    assert doc.metadata.extra["sheet_names"] == ["成績", "備註"]
    assert doc.metadata.extra["formula_count"] == 1


def test_sheets_and_used_range(sample: Sample) -> None:
    doc = XlsxHandler().open(sample("xlsx"))
    sheets = doc.sheets()
    assert [s.name for s in sheets] == ["成績", "備註"]
    assert sheets[0].dimensions == "A1:E4"
    assert (sheets[0].max_row, sheets[0].max_column) == (4, 5)
    assert sheets[1].dimensions == "A1:A1"
    assert [(p.kind, p.index, p.name) for p in doc.parts()] == [
        ("sheet", 0, "成績"),
        ("sheet", 1, "備註"),
    ]


def test_sheet_state_and_empty_sheet(tmp_path: Path) -> None:
    doc = XlsxHandler().open(_rich_xlsx(tmp_path))
    states = {s.name: s.state for s in doc.sheets()}
    assert states == {"資料": "visible", "隱藏": "hidden", "空白": "visible"}
    empty = doc.sheets()[2]
    assert empty.dimensions is None
    assert (empty.max_row, empty.max_column) == (0, 0)
    assert doc.rows("空白") == []
    assert doc.cells("空白") == []


def test_cell_values_and_types(sample: Sample) -> None:
    doc = XlsxHandler().open(sample("xlsx"))
    cells = {c.coordinate: c for c in doc.cells("成績")}
    assert cells["A1"].value == TABLE_HEADER[0]
    assert cells["A1"].type == "string"
    assert cells["C2"].value == TABLE_ROWS[0][2]
    assert cells["C2"].type == "number"
    assert "D2" not in cells
    assert (cells["C2"].row, cells["C2"].column) == (2, 3)


def test_value_types_date_bool_error(tmp_path: Path) -> None:
    doc = XlsxHandler().open(_rich_xlsx(tmp_path))
    cells = {c.coordinate: c for c in doc.cells("資料")}
    assert cells["B2"].type == "boolean" and cells["B2"].value is True
    assert cells["C2"].type == "date"
    assert cells["D2"].type == "date" and cells["D2"].value == datetime(2026, 1, 2, 3, 4, 5)
    assert cells["B3"].type == "error"


def test_formula_without_cached_value(sample: Sample) -> None:
    doc = XlsxHandler().open(sample("xlsx"))
    cell = next(c for c in doc.cells("成績") if c.coordinate == "E2")
    assert cell.type == "formula"
    assert cell.formula == "=AVERAGE(C2:C4)"
    assert cell.value is None


def test_formula_with_cached_value(tmp_path: Path) -> None:
    doc = XlsxHandler().open(_rich_xlsx(tmp_path))
    cell = next(c for c in doc.cells("資料") if c.coordinate == "A3")
    assert cell.formula == "=SUM(A2:A2)"
    assert cell.value == 12.5
    assert doc.rows("資料")[2][0] == 12.5


def test_merged_cells(tmp_path: Path) -> None:
    doc = XlsxHandler().open(_rich_xlsx(tmp_path))
    assert doc.merged_cells("資料") == ["A1:C1"]
    assert doc.sheets()[0].merged_ranges == ["A1:C1"]
    assert doc.merged_cells(1) == []


def test_style_summary(tmp_path: Path) -> None:
    doc = XlsxHandler().open(_rich_xlsx(tmp_path))
    cells = {c.coordinate: c for c in doc.cells("資料")}
    a1 = cells["A1"]
    assert (a1.bold, a1.italic, a1.font_name, a1.font_size) == (True, True, "Arial", 14.0)
    assert a1.fill_color == "00FFFF00"
    assert cells["A2"].number_format == "0.00"
    assert cells["B2"].fill_color is None
    summary = doc.style_summary("資料")
    assert summary["bold_cells"] == 1
    assert summary["italic_cells"] == 1
    assert summary["number_formats"]["0.00"] == 1
    assert summary["fonts"]["Arial"] == 1
    assert summary["fills"] == {"00FFFF00": 1}


def test_defined_names(tmp_path: Path) -> None:
    doc = XlsxHandler().open(_rich_xlsx(tmp_path))
    names = {(n.name, n.value, n.scope) for n in doc.defined_names()}
    assert names == {("Total", "資料!$A$3", None), ("Local", "隱藏!$A$1", "隱藏")}
    assert doc.metadata.extra["defined_names"] == ["Total", "Local"]


def test_core_properties(tmp_path: Path) -> None:
    doc = XlsxHandler().open(_rich_xlsx(tmp_path))
    assert doc.metadata.title == "豐富活頁簿"
    assert doc.metadata.author == "劉老師"


def test_sheet_lookup_by_name_and_index(sample: Sample) -> None:
    doc = XlsxHandler().open(sample("xlsx"))
    assert doc.rows("備註") == doc.rows(1) == [[BODY]]
    with pytest.raises(KeyError):
        doc.rows("nope")
    with pytest.raises(KeyError):
        doc.cells(9)


def test_text_in_sheet_order(sample: Sample) -> None:
    text = XlsxHandler().open(sample("xlsx")).text()
    lines = text.splitlines()
    assert lines[0] == "成績"
    assert lines[1] == "\t".join([*TABLE_HEADER, "", "平均"])
    assert "王小明\t研發部\t90" in text
    assert text.index("成績") < text.index("備註")
    assert text.endswith(BODY)


def test_text_formats_bool_and_date(tmp_path: Path) -> None:
    text = XlsxHandler().open(_rich_xlsx(tmp_path)).text()
    assert "TRUE" in text
    assert "2026-01-02" in text
    assert "隱藏\nsecret" in text


def test_open_from_bytes(sample: Sample) -> None:
    doc = XlsxHandler().open(sample("xlsx").read_bytes())
    assert doc.path is None
    assert doc.sheets()[0].name == "成績"
    with pytest.raises(ValueError):
        doc.save()


def test_save_roundtrip_preserves_formulas(sample: Sample, tmp_path: Path) -> None:
    doc = XlsxHandler().open(sample("xlsx"))
    out = doc.save(tmp_path / "out.xlsx")
    again = XlsxHandler().open(out)
    cell = next(c for c in again.cells("成績") if c.coordinate == "E2")
    assert cell.formula == "=AVERAGE(C2:C4)"
    assert again.rows("成績")[1][:3] == list(TABLE_ROWS[0])
    again.close()
    doc.close()


def test_save_in_place(sample: Sample, tmp_path: Path) -> None:
    path = tmp_path / "copy.xlsx"
    path.write_bytes(sample("xlsx").read_bytes())
    doc = XlsxHandler().open(path)
    assert doc.save() == path


@pytest.mark.parametrize("data", [b"", b"not a zip", b"PK\x03\x04garbage"])
def test_corrupt_input(data: bytes) -> None:
    with pytest.raises(CorruptFileError):
        XlsxHandler().open(data)


def test_zip_without_workbook_part() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(CorruptFileError):
        XlsxHandler().open(buf.getvalue())


def test_missing_file(tmp_path: Path) -> None:
    with pytest.raises(CorruptFileError):
        XlsxHandler().open(tmp_path / "missing.xlsx")


def test_cli_info(sample: Sample) -> None:
    from typer.testing import CliRunner

    from anydoc.cli import app

    result = CliRunner().invoke(app, ["info", str(sample("xlsx"))])
    assert result.exit_code == 0, result.output
    assert "成績" in result.output
