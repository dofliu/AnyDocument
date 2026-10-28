"""Tests for the DOCX reader (T012)."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document as NewDocx
from docx.enum.text import WD_BREAK

import anydoc
from anydoc.core.errors import CorruptFileError
from anydoc.core.model import Capability, Document
from anydoc.core.registry import registry
from anydoc.formats.docx import DocxDocument, DocxHandler
from tests.fixtures.generate import BODY, TABLE_HEADER, TABLE_ROWS, TITLE, make_png

Sample = Callable[[str], Path]

_FOOTNOTES_XML = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:id="1"><w:p><w:r><w:t>註腳一：中文說明</w:t></w:r></w:p></w:footnote>'
    '<w:footnote w:id="2"><w:p><w:r><w:t>Second note</w:t></w:r></w:p></w:footnote>'
    "</w:footnotes>"
)
_PNG = make_png()


@pytest.fixture(autouse=True)
def _handler() -> None:
    registry.register(DocxHandler, replace=True)


def _add_footnotes(data: bytes) -> bytes:
    """Inject a footnotes part into a DOCX package."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            blob = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                blob = blob.replace(
                    b"</Types>",
                    b'<Override PartName="/word/footnotes.xml" ContentType="application/'
                    b'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
                )
            elif item.filename == "word/_rels/document.xml.rels":
                blob = blob.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/'
                    b'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
                    b"</Relationships>",
                )
            zout.writestr(item, blob)
        zout.writestr("word/footnotes.xml", _FOOTNOTES_XML)
    return out.getvalue()


def _rich_docx(tmp_path: Path) -> Path:
    doc = NewDocx()
    doc.core_properties.title = "豐富文件"
    doc.core_properties.author = "劉老師"
    doc.core_properties.subject = "主旨"
    doc.core_properties.created = datetime(2026, 1, 2, 3, 4, 5)
    doc.add_paragraph("文件標題", style="Title")
    doc.add_heading("第一章", level=1)
    doc.add_paragraph("第一段內容。")
    doc.add_heading("第一節", level=2)
    table = doc.add_table(rows=2, cols=2)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.text = f"儲存格{i}{j}"
    doc.add_paragraph("表格之後。")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    pic = tmp_path / "pic.png"
    pic.write_bytes(_PNG)
    doc.add_picture(str(pic))
    doc.sections[0].header.paragraphs[0].text = "頁首文字"
    doc.sections[0].footer.paragraphs[0].text = "頁尾文字"
    buf = io.BytesIO()
    doc.save(buf)
    path = tmp_path / "rich.docx"
    path.write_bytes(_add_footnotes(buf.getvalue()))
    return path


def test_entry_point_and_capabilities() -> None:
    from anydoc.core.registry import HandlerRegistry

    reg = HandlerRegistry()
    reg.load_entry_points()
    assert reg.get("docx") is DocxHandler
    assert Capability.READ in DocxHandler.capabilities
    assert Capability.EXTRACT_TEXT in DocxHandler.capabilities


def test_sniff() -> None:
    assert DocxHandler.sniff(b"", Path("a.DOCX")) == 0.9
    assert DocxHandler.sniff(b"", Path("a.doc")) == 0.0
    assert DocxHandler.sniff(b"", None) == 0.0


def test_open_sample_via_public_api(sample: Sample) -> None:
    doc = anydoc.open(sample("docx"))
    assert isinstance(doc, DocxDocument)
    assert isinstance(doc, Document)
    assert doc.format == "docx"
    assert doc.metadata.title == TITLE
    assert doc.metadata.author == "AnyDocument"
    assert doc.metadata.created is None or isinstance(doc.metadata.created, datetime)


def test_paragraph_styles_and_heading_levels(tmp_path: Path) -> None:
    doc = DocxHandler().open(_rich_docx(tmp_path))
    by_text = {p.text: p for p in doc.paragraphs()}
    assert by_text["文件標題"].style == "Title"
    assert by_text["文件標題"].heading_level == 0
    assert by_text["第一章"].heading_level == 1
    assert by_text["第一節"].heading_level == 2
    assert by_text["第一段內容。"].heading_level is None
    assert [(h.heading_level, h.text) for h in doc.headings()] == [
        (0, "文件標題"),
        (1, "第一章"),
        (2, "第一節"),
    ]
    assert doc.metadata.extra["headings"][1] == {"level": 1, "title": "第一章"}


def test_sample_table(sample: Sample) -> None:
    doc = DocxHandler().open(sample("docx"))
    [table] = doc.tables()
    assert table.rows[0] == list(TABLE_HEADER)
    assert table.rows[1:] == [[str(c) for c in r] for r in TABLE_ROWS]
    assert doc.metadata.extra["table_count"] == 1


def test_images_listed(tmp_path: Path) -> None:
    doc = DocxHandler().open(_rich_docx(tmp_path))
    [image] = doc.images()
    assert image.content_type == "image/png"
    assert image.size == len(_PNG)
    assert image.name.endswith(".png")
    assert doc.metadata.extra["image_count"] == 1


def test_no_images(sample: Sample) -> None:
    assert DocxHandler().open(sample("docx")).images() == []


def test_headers_footers(tmp_path: Path) -> None:
    doc = DocxHandler().open(_rich_docx(tmp_path))
    assert doc.headers_footers() == [{"header": "頁首文字", "footer": "頁尾文字"}]


def test_footnotes(tmp_path: Path) -> None:
    doc = DocxHandler().open(_rich_docx(tmp_path))
    assert doc.footnotes() == ["註腳一：中文說明", "Second note"]
    assert doc.metadata.extra["footnote_count"] == 2


def test_no_footnotes(sample: Sample) -> None:
    assert DocxHandler().open(sample("docx")).footnotes() == []


def test_core_properties(tmp_path: Path) -> None:
    doc = DocxHandler().open(_rich_docx(tmp_path))
    assert doc.metadata.title == "豐富文件"
    assert doc.metadata.author == "劉老師"
    assert doc.metadata.extra["subject"] == "主旨"
    assert doc.metadata.created == datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC)
    assert doc.metadata.to_dict()["created"] == "2026-01-02T03:04:05+00:00"


def test_text_preserves_reading_order(tmp_path: Path) -> None:
    text = DocxHandler().open(_rich_docx(tmp_path)).text()
    order = [
        "文件標題",
        "第一章",
        "第一段內容。",
        "第一節",
        "儲存格00\t儲存格01",
        "儲存格10\t儲存格11",
        "表格之後。",
    ]
    positions = [text.index(s) for s in order]
    assert positions == sorted(positions)


def test_sample_text_order(sample: Sample) -> None:
    text = DocxHandler().open(sample("docx")).text()
    assert text.index(TITLE) < text.index(BODY) < text.index("\t".join(TABLE_HEADER))
    assert "王小明\t研發部\t90" in text


def test_parts_sections_with_page_size(sample: Sample) -> None:
    [part] = DocxHandler().open(sample("docx")).parts()
    assert part.kind == "section"
    assert part.size is not None
    assert part.size[0] > 0
    assert part.size[1] > part.size[0]  # portrait


def test_open_from_bytes(sample: Sample) -> None:
    doc = DocxHandler().open(sample("docx").read_bytes())
    assert doc.path is None
    assert TITLE in doc.text()


def test_save_roundtrip(sample: Sample, tmp_path: Path) -> None:
    doc = DocxHandler().open(sample("docx"))
    out = doc.save(tmp_path / "copy.docx")
    reopened = DocxHandler().open(out)
    assert reopened.text() == doc.text()
    assert reopened.metadata.title == TITLE


def test_save_without_path_raises(sample: Sample) -> None:
    with pytest.raises(ValueError, match="No path"):
        DocxHandler().open(sample("docx").read_bytes()).save()


def test_empty_document() -> None:
    buf = io.BytesIO()
    NewDocx().save(buf)
    doc = DocxHandler().open(buf.getvalue())
    assert doc.text() == ""
    assert doc.tables() == []
    assert doc.footnotes() == []
    assert doc.metadata.title is None


@pytest.mark.parametrize("data", [b"", b"not a zip", b"PK\x03\x04garbage"])
def test_corrupt_input(data: bytes) -> None:
    with pytest.raises(CorruptFileError):
        DocxHandler().open(data)


def test_zip_without_document_part() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(CorruptFileError):
        DocxHandler().open(buf.getvalue())


def test_missing_file(tmp_path: Path) -> None:
    with pytest.raises(CorruptFileError):
        DocxHandler().open(tmp_path / "nope.docx")


def test_invalid_footnotes_part(tmp_path: Path) -> None:
    data = _rich_docx(tmp_path).read_bytes().replace(b"<w:footnotes", b"<w:footnotes <<")
    # footnotes.xml is stored (not compressed) so the replacement hits the part
    with pytest.raises(CorruptFileError):
        DocxHandler().open(data)


def test_info_via_cli(sample: Sample) -> None:
    import json

    from typer.testing import CliRunner

    from anydoc.cli import app

    result = CliRunner().invoke(app, ["info", str(sample("docx"))])
    assert result.exit_code == 0, result.output
    assert json.loads(result.output)["metadata"]["title"] == TITLE
