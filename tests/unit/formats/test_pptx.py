"""Tests for the PPTX reader (T014)."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

import anydoc
from anydoc.core.errors import CorruptFileError
from anydoc.core.model import Capability, Document
from anydoc.core.registry import registry
from anydoc.formats.pptx import PptxDocument, PptxHandler
from tests.fixtures.generate import BODY, TITLE

Sample = Callable[[str], Path]


@pytest.fixture(autouse=True)
def _handler() -> None:
    registry.register(PptxHandler, replace=True)


def _rich_pptx(tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = "圖形頁"  # type: ignore[union-attr]
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    box.text_frame.text = "文字框內容"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    for r, row in enumerate([["姓名", "分數"], ["小明", "95"]]):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    data = CategoryChartData()
    data.categories = ["甲", "乙"]
    data.add_series("S1", (1, 2))
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(2), Inches(3), Inches(3), data
    )
    chart_shape.chart.has_title = True
    chart_shape.chart.chart_title.text_frame.text = "圖表標題"
    png = io.BytesIO()
    from PIL import Image

    Image.new("RGB", (4, 4), "red").save(png, format="PNG")
    png.seek(0)
    slide.shapes.add_picture(png, Inches(0), Inches(0), Inches(1), Inches(1))
    group = slide.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(1))
    inner.text_frame.text = "群組內文字"
    path = tmp_path / "rich.pptx"
    prs.core_properties.title = "簡報標題"
    prs.core_properties.author = "作者"
    prs.save(str(path))
    return path


def test_entry_point_and_capabilities() -> None:
    from importlib.metadata import entry_points

    eps = {e.name: e for e in entry_points(group="anydoc.formats")}
    assert eps["pptx"].load() is PptxHandler
    assert Capability.READ in PptxHandler.capabilities
    assert Capability.EXTRACT_TEXT in PptxHandler.capabilities


def test_sniff() -> None:
    assert PptxHandler.sniff(b"", Path("a.PPTX")) > 0.5
    assert PptxHandler.sniff(b"", Path("a.docx")) == 0.0
    assert PptxHandler.sniff(b"", None) == 0.0


def test_open_sample_via_public_api(sample: Sample) -> None:
    doc = anydoc.open(sample("pptx"))
    assert isinstance(doc, Document)
    assert isinstance(doc, PptxDocument)
    assert doc.metadata.format == "pptx"
    assert doc.metadata.page_count == 2


def test_slides_layouts_titles_notes(sample: Sample) -> None:
    doc = PptxHandler().open(sample("pptx"))
    s1, s2 = doc.slides()
    assert s1.layout == "Title Slide"
    assert s1.title == TITLE
    assert s2.layout == "Title and Content"
    assert s2.title == "第二頁"
    assert s1.notes == ""
    assert s2.notes == "講者備忘稿"
    assert "Title Slide" in doc.layouts()
    assert doc.slide(1).title == "第二頁"
    with pytest.raises(IndexError):
        doc.slide(5)
    with pytest.raises(IndexError):
        doc.slide(-1)


def test_parts(sample: Sample) -> None:
    parts = PptxHandler().open(sample("pptx")).parts()
    assert [(p.kind, p.index, p.name) for p in parts] == [
        ("slide", 0, TITLE),
        ("slide", 1, "第二頁"),
    ]


def test_part_name_fallback_without_title() -> None:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])  # blank
    buf = io.BytesIO()
    prs.save(buf)
    doc = PptxHandler().open(buf.getvalue())
    assert doc.slides()[0].title is None
    assert doc.parts()[0].name == "Slide 1"


def test_shape_tree_kinds(tmp_path: Path) -> None:
    doc = PptxHandler().open(_rich_pptx(tmp_path))
    shapes = {s.kind: s for s in doc.slide(0).shapes}
    assert set(shapes) >= {"text", "table", "chart", "picture", "group"}
    assert shapes["table"].table == [["姓名", "分數"], ["小明", "95"]]
    assert shapes["chart"].chart_type is not None
    assert "COLUMN_CLUSTERED" in shapes["chart"].chart_type
    assert shapes["chart"].text == "圖表標題"
    assert [c.text for c in shapes["group"].children] == ["群組內文字"]
    textbox = next(s for s in doc.slide(0).shapes if s.text == "文字框內容")
    assert textbox.left == Inches(1)
    assert textbox.width == Inches(3)
    assert not textbox.placeholder
    assert doc.slide(0).shapes[0].placeholder  # the title placeholder


def test_slide_text_flattens_groups(tmp_path: Path) -> None:
    doc = PptxHandler().open(_rich_pptx(tmp_path))
    text = doc.slide_text(0)
    for needle in ("圖形頁", "文字框內容", "姓名\t分數", "圖表標題", "群組內文字"):
        assert needle in text


def test_text_in_reading_order_with_notes(sample: Sample) -> None:
    text = PptxHandler().open(sample("pptx")).text()
    order = [text.index(x) for x in (TITLE, BODY, "第二頁", "重點一", "重點二", "講者備忘稿")]
    assert order == sorted(order)


def test_core_properties(tmp_path: Path) -> None:
    meta = PptxHandler().open(_rich_pptx(tmp_path)).metadata
    assert meta.title == "簡報標題"
    assert meta.author == "作者"
    assert meta.extra["slide_count"] == 1
    assert meta.extra["slide_width"] > 0
    assert "Title Slide" in meta.extra["layouts"]


def test_open_from_bytes(sample: Sample) -> None:
    doc = PptxHandler().open(sample("pptx").read_bytes())
    assert doc.path is None
    assert len(doc.slides()) == 2


def test_save_roundtrip(sample: Sample, tmp_path: Path) -> None:
    doc = PptxHandler().open(sample("pptx"))
    out = doc.save(tmp_path / "copy.pptx")
    again = PptxHandler().open(out)
    assert again.text() == doc.text()


def test_save_in_place(sample: Sample, tmp_path: Path) -> None:
    target = tmp_path / "inplace.pptx"
    target.write_bytes(sample("pptx").read_bytes())
    doc = PptxHandler().open(target)
    assert doc.save() == target
    doc.close()


def test_save_without_path_raises(sample: Sample) -> None:
    doc = PptxHandler().open(sample("pptx").read_bytes())
    with pytest.raises(ValueError):
        doc.save()


@pytest.mark.parametrize("data", [b"", b"not a zip", b"PK\x03\x04garbage"])
def test_corrupt_input(data: bytes) -> None:
    with pytest.raises(CorruptFileError):
        PptxHandler().open(data)


def test_zip_without_presentation_part() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(CorruptFileError):
        PptxHandler().open(buf.getvalue())


def test_missing_file(tmp_path: Path) -> None:
    with pytest.raises(CorruptFileError):
        PptxHandler().open(tmp_path / "nope.pptx")


def test_cli_info(sample: Sample) -> None:
    from typer.testing import CliRunner

    from anydoc.cli import app

    result = CliRunner().invoke(app, ["info", str(sample("pptx"))])
    assert result.exit_code == 0, result.output
    assert "第二頁" in result.output
