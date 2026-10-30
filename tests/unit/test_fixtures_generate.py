"""Tests for the sample generator: every sample opens with its native library."""

from __future__ import annotations

import io
import json
from collections.abc import Callable
from pathlib import Path

import pytest

from anydoc import detect
from tests.fixtures import generate

Sample = Callable[[str], Path]


def test_all_samples_generated(sample: Sample) -> None:
    for fmt in generate.SAMPLES:
        path = sample(fmt)
        assert path.is_file()
        assert 0 < path.stat().st_size < 200_000


def test_unknown_sample_raises(sample: Sample) -> None:
    with pytest.raises(KeyError):
        sample("nope")
    with pytest.raises(KeyError):
        generate.generate_sample("nope")


def test_generate_sample_to_custom_dir(tmp_path: Path) -> None:
    path = generate.generate_sample("txt", tmp_path / "sub")
    assert path.parent == tmp_path / "sub"
    assert generate.TITLE in path.read_text(encoding="utf-8")


def test_txt_json(sample: Sample) -> None:
    assert generate.BODY in sample("txt").read_text(encoding="utf-8")
    data = json.loads(sample("json").read_text(encoding="utf-8"))
    assert data["title"] == generate.TITLE


def test_csv_encodings(sample: Sample) -> None:
    assert sample("csv").read_text(encoding="utf-8").startswith("姓名,")
    raw_bom = sample("csv-bom").read_bytes()
    assert raw_bom.startswith(b"\xef\xbb\xbf")
    assert raw_bom.decode("utf-8-sig").startswith("姓名,")
    raw_big5 = sample("csv-big5").read_bytes()
    with pytest.raises(UnicodeDecodeError):
        raw_big5.decode("utf-8")
    assert raw_big5.decode("big5").startswith("姓名,")


def test_images(sample: Sample) -> None:
    from PIL import Image

    with Image.open(sample("png")) as png:
        assert png.format == "PNG" and png.size == (64, 48)
    with Image.open(sample("jpg")) as jpg:
        assert jpg.format == "JPEG"


def test_docx(sample: Sample) -> None:
    from docx import Document

    doc = Document(str(sample("docx")))
    texts = [p.text for p in doc.paragraphs]
    assert generate.TITLE in texts and generate.BODY in texts
    assert doc.tables[0].rows[1].cells[0].text == "王小明"
    assert doc.core_properties.title == generate.TITLE


def test_xlsx(sample: Sample) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(sample("xlsx"))
    assert wb.sheetnames == ["成績", "備註"]
    assert wb["成績"]["A2"].value == "王小明"
    assert wb["成績"]["E2"].value == "=AVERAGE(C2:C4)"


def test_pptx(sample: Sample) -> None:
    from pptx import Presentation

    prs = Presentation(str(sample("pptx")))
    assert len(prs.slides) == 2
    assert prs.slides[0].shapes.title is not None
    assert prs.slides[0].shapes.title.text == generate.TITLE
    assert prs.slides[1].notes_slide.notes_text_frame is not None
    assert prs.slides[1].notes_slide.notes_text_frame.text == "講者備忘稿"


def test_pdf(sample: Sample) -> None:
    import pdfplumber
    from pypdf import PdfReader

    reader = PdfReader(str(sample("pdf")))
    assert len(reader.pages) == 2
    with pdfplumber.open(sample("pdf")) as pdf:
        assert "hello world" in (pdf.pages[1].extract_text() or "")


def test_dxf(sample: Sample) -> None:
    import ezdxf

    doc = ezdxf.readfile(sample("dxf"))
    assert {"牆", "標註"} <= {layer.dxf.name for layer in doc.layers}
    kinds = sorted(e.dxftype() for e in doc.modelspace())
    assert kinds == ["CIRCLE", "LINE", "TEXT"]
    texts = [e.dxf.text for e in doc.modelspace().query("TEXT")]
    assert texts == [generate.TITLE]


@pytest.mark.parametrize(
    ("fmt", "expected"),
    [
        ("docx", "docx"),
        ("xlsx", "xlsx"),
        ("pptx", "pptx"),
        ("pdf", "pdf"),
        ("dxf", "dxf"),
        ("png", "png"),
        ("jpg", "jpeg"),
        ("csv", "csv"),
        ("csv-bom", "csv"),
        ("csv-big5", "csv"),
        ("json", "json"),
        ("txt", "txt"),
        ("rtf", "rtf"),
        ("html", "html"),
        ("epub", "epub"),
    ],
)
def test_samples_detected(sample: Sample, fmt: str, expected: str) -> None:
    assert detect(sample(fmt)).format == expected


def test_pdf_bytes_roundtrip() -> None:
    from pypdf import PdfReader

    assert len(PdfReader(io.BytesIO(generate.make_pdf())).pages) == 2
